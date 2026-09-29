import openpyxl as ox
from DB_Feature_Settings import Add_Feature_Headers, Clean_Action
import Player_Settings as P
import regex as re


#GLOBALS
Turn_Player = ""
Turn_Count = 0
P1Deck = 40
P2Deck = 41
# Blue (0) Red(1)
Went_First = 0

#Create Players
P1 = P.Player()
P2 = P.Player()

P1.set_Deck_Size(P1Deck)
P2.set_Deck_Size(P2Deck)

# Read the Dueling Book Log Excel
wb = ox.load_workbook("Db_Test.xlsx")
ws = wb.worksheets[0]

#This worksheet will have the setup with all the data
wb.create_sheet(title='Db_Log', index=None)
ws1 = wb['Db_Log']
ws1.append(Add_Feature_Headers())

# Get DB Log Action Length
Log_length = int(len(ws['A:A']))

# The Check color of the Text in cell, Looking for:
# Blue: FF0000FF
# Red:  FFFF0000

# Add what Player is each action
for row in range(1,Log_length+1):
    r = str(row)
    Check_row = 'A' + r
    #current_row = 'B' + r

    #Value of the Action (Cell)
    Action = Clean_Action(ws[Check_row].value)

    #Move Action to Db_Log Sheet
    ws1.cell(row + 1, 1, value=Action)

    #Color of the Action (Cell)
    Player = ws[Check_row].font.color.rgb

    #Turn Count
    ws1.cell(row + 1, 2, value=Turn_Count)

    # Mark 1 for Player Action, 0 for not there Action based on Text Color
    #Blue (Player 1)
    if Player == 'FF0000FF':
        #ws[current_row] = 'Player1'
        ws1.cell(row + 1, 4, value=1)
        ws1.cell(row + 1, 5, value=0)

        # Find out who is Turn Player then alternate from there
        if Action.lower() == 'Chose to go first'.lower():
            Turn_Player = 'Player1'

        # Game State Actions:

        # Draw a card from Deck
        if Action.lower() == 'Drew a card'.lower():
            P1.Increase_Hand_Size(1)
            P1.Decrease_Deck_Size(1)

        # Life Points
            # Increase LP
        if re.search("^Gained.*LP$",Action):
            LP =int(Action[7:11])
            P1.Increase_Life_Points(LP)

            # Decrease LP
        if re.search("^Lost.*LP$",Action):
            LP =int(Action[5:9])
            P1.Decrease_Life_Points(LP)

            # Set that this player is going 1st
        if Action.lower() == 'Chose to go first'.lower():
            P1.set_Going_First()





#-------------------------------------------------------------------------------------------------#

    #Red (Player 2)
    if Player == 'FFFF0000':
        #ws[current_row] = 'Player2'
        ws1.cell(row + 1, 4, value=0)
        ws1.cell(row + 1, 5, value=1)

        #Turn Player
        if Action.lower() == 'Chose to go first'.lower():
            Turn_Player = 'Player2'

        # Game State Actions:
            # Draw a card from Deck
        if Action.lower() == 'Drew a card'.lower():
            P2.Increase_Hand_Size(1)
            P2.Decrease_Deck_Size(1)

        # Life Points
            #Increase LP
        if re.search("^Gained.*LP$",Action):
            LP =int(Action[7:11])
            P2.Increase_Life_Points(LP)

            #Decrease LP
        if re.search("^Lost.*LP$",Action):
            LP =int(Action[5:9])
            P2.Decrease_Life_Points(LP)

            # Set that this player is going 1st
        if Action.lower() == 'Chose to go first'.lower():
            P2.set_Going_First()

#----------------------------------------------------------------------------------------------#

    #Keep track of Turn Player until the end there turn
    if Turn_Player == 'Player1':
        ws1.cell(row + 1, 6, value=1)
        ws1.cell(row + 1, 7, value=0)
    else:
        ws1.cell(row + 1, 6, value=0)
        ws1.cell(row + 1, 7, value=1)

    if Action.lower() == 'Ended turn'.lower():
        if Turn_Player == 'Player1':
            Turn_Player = 'Player2'
            Turn_Count += 1
        else:
            Turn_Player = 'Player1'
            Turn_Count += 1

            # Reset Stats for new match
    if Action.lower() == 'Admitted defeat'.lower():
        P1.new_match(P1Deck)
        P2.new_match(P2Deck)


    #Current Game State
        # Hand Size
    ws1.cell(row + 1, 8, value=P1.Hand_Size)
    ws1.cell(row + 1, 9, value=P2.Hand_Size)

        # Deck Size
    ws1.cell(row + 1, 10, value=P1.Deck_Size)
    ws1.cell(row + 1, 11, value=P2.Deck_Size)

        # Life Points
    ws1.cell(row + 1, 26, value=P1.Life_Points)
    ws1.cell(row + 1, 27, value=P2.Life_Points)

        # Going First
    ws1.cell(row + 1, 28, value=P1.Went_First)
    ws1.cell(row + 1, 29, value=P2.Went_First)

# Save the file
wb.save("Db_Test.xlsx")
print("DONE")
